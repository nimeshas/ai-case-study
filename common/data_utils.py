"""Shared, local-only loading and validation. Source datasets are never rewritten."""
import csv
import hashlib
import json
from pathlib import Path

import openpyxl
from common.paths import DATA, RESULTS

from scoring.score_word_counts import read_csv_robust

COMPONENTS = ['initiative', 'resolution', 'recognition']
METRIC_COLUMNS = [*COMPONENTS, 'agency_score', 'competence_word_count', 'emotional_support_word_count']
RESPONSE_COLUMNS = ['response_id', 'prompt_id', 'occupation', 'gender_version', 'model',
                    'repetition', 'response_text', *METRIC_COLUMNS]
SCORE_COLUMNS = ['response_id', *COMPONENTS, 'agency_score', 'explanation']
SOURCES = [('responses_chatgpt.csv', 'ChatGPT', 'gpt'),
           ('responses_gemini.csv', 'Gemini', 'gem')]


def write_csv(path, rows, columns):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def story_hash(text):
    return hashlib.sha256(text.encode('utf-8')).hexdigest()


def load_responses(root=DATA, output=RESULTS):
    """Validate the cleaned 90-per-model datasets. Never alter exclusion evidence."""
    root, output = Path(root), Path(output)
    output.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(root / 'prompt_dataset.xlsx', read_only=True, data_only=True)
    prompts, cells, errors, records = {}, set(), [], []
    for values in wb['Prompts'].iter_rows(min_row=2, values_only=True):
        pid, occupation, _, gender, _ = values[:5]
        if pid is None:
            continue
        if pid in prompts or (occupation, gender) in cells:
            errors.append(f'Duplicate prompt ID or occupation/gender: {pid}')
        prompts[pid] = (occupation, gender)
        cells.add((occupation, gender))
    wb.close()
    occupations = {o for o, g in prompts.values()}
    if len(prompts) != 30 or len(occupations) != 10 or cells != {
        (o, g) for o in occupations for g in ['Neutral', 'Male', 'Female']
    }:
        errors.append('Prompt workbook must contain 10 occupations x 3 gender conditions.')
    counts = {}
    for filename, model, suffix in SOURCES:
        fields, rows = read_csv_robust(root / filename)
        if fields is None or len(fields) != len(set(fields)):
            raise ValueError(f'{filename}: missing or duplicate column names')
        required = RESPONSE_COLUMNS
        if not set(required).issubset(fields or []):
            raise ValueError(f'{filename}: missing columns {set(required) - set(fields or [])}')
        if fields[-6:] != METRIC_COLUMNS or '' in fields or 'date_generated' in fields:
            errors.append(f'{filename}: response column order/schema is incorrect')
        expected = {f'R_{pid}_{suffix}_{rep}' for pid in prompts for rep in range(1, 4)}
        seen = set()
        for row_number, row in enumerate(rows, 2):
            rid, pid = row['response_id'], row['prompt_id']
            if None in row or any(row[c] is None for c in fields):
                errors.append(f'{filename} record {row_number}: malformed CSV row')
            if rid in seen:
                errors.append(f'{filename} record {row_number}: duplicate ID {rid}; no records silently removed')
            seen.add(rid)
            if (rid not in expected or pid not in prompts or row['model'] != model
                or row['repetition'] not in ['1', '2', '3']):
                errors.append(f'{filename} record {row_number}: invalid ID/model/repetition')
            elif (rid != f"R_{pid}_{suffix}_{row['repetition']}"
                  or (row['occupation'], row['gender_version']) != prompts[pid]):
                errors.append(f'{filename} record {row_number}: metadata differs from prompt workbook')
            text = row['response_text'] or ''
            if not text.strip() or text.strip().startswith('[ERROR:'):
                errors.append(f'{filename} record {row_number}: empty/error story')
            if any(row.get(c) not in ('', None) for c in COMPONENTS + ['agency_score']):
                try:
                    validate_scores(row)
                except ValueError as error:
                    errors.append(f'{filename} record {row_number}: {error}')
            records.append(row)
        if expected - seen:
            errors.append(f'{filename}: missing IDs {sorted(expected - seen)}')
        if len(rows) != 90 or len(seen) != 90:
            errors.append(f'{filename}: expected 90 rows and unique IDs; found {len(rows)} rows, {len(seen)} IDs')
        counts[model] = dict(rows=len(rows), unique_ids=len(seen))
    report = dict(source_sha256={name: hashlib.sha256((root / name).read_bytes()).hexdigest()
                                for name, _, _ in SOURCES},
                  records=len(records), counts_by_model=counts, errors=errors,
                  original_backups='data/archive/original', historical_exclusions='results/excluded_duplicates.csv')
    (output / 'validation.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    if errors:
        raise ValueError('Dataset validation failed; see ' + str(output / 'validation.json'))
    return records


def validate_scores(row):
    """Validate binary values and total; accept Excel's numeric 0.0 and 1.0."""
    result = dict(row)
    for column in COMPONENTS + ['agency_score']:
        value = row.get(column)
        if isinstance(value, bool) or value is None or str(value).strip() == '':
            raise ValueError(f"{row.get('response_id')}: missing/invalid {column}")
        try:
            number = float(value)
        except (ValueError, TypeError):
            raise ValueError(f"{row.get('response_id')}: invalid {column}") from None
        allowed = [0, 1, 2, 3] if column == 'agency_score' else [0, 1]
        if number not in allowed:
            raise ValueError(f"{row.get('response_id')}: invalid {column}")
        result[column] = int(number)
    if result['agency_score'] != sum(result[c] for c in COMPONENTS):
        raise ValueError(f"{row.get('response_id')}: total does not equal component sum")
    return result


def load_scores(path, sheet=None, id_column='response_id',
                component_columns=None, total_column=None, explanation_column=None):
    """Explicit column mapping supports existing manual workbook conventions."""
    path = Path(path)
    if path.suffix.lower() == '.xlsx':
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.active
        values = iter(ws.values)
        fields = next(values)
        rows = [dict(zip(fields, values)) for values in values]
        wb.close()
    else:
        fields, rows = read_csv_robust(path)
    if fields is None or len(fields) != len(set(fields)):
        raise ValueError(f'{path}: missing or duplicate column names')

    component_columns = component_columns or [c if c in fields else f'{c} (0/1)' for c in COMPONENTS]
    total_column = total_column or ('agency_score' if 'agency_score' in fields else 'agency_score (0-3)')
    explanation_column = explanation_column or next(
        (c for c in ['explanation', 'judge_reason', 'evidence / notes'] if c in fields), 'explanation')
    mapping = dict(zip(COMPONENTS, component_columns))
    mapping.update(response_id=id_column, agency_score=total_column)
    missing = set(mapping.values()) - set(fields)
    if missing:
        raise ValueError(f'{path}: missing columns {sorted(missing)}; use column-mapping options')
    result, seen = [], set()
    for raw in rows:
        if all(v is None or str(v).strip() == '' for v in raw.values()):
            continue
        row = {key: raw[value] for key, value in mapping.items()}
        if row['response_id'] is None or not str(row['response_id']).strip():
            raise ValueError(f'{path}: scored row has no Story ID')
        if row['response_id'] in seen:
            raise ValueError(f"{path}: duplicate score ID {row['response_id']}")
        seen.add(row['response_id'])
        # An unlabelled template row is not a zero score. Partial labels are errors.
        if all(row[c] is None or str(row[c]).strip() == '' for c in COMPONENTS + ['agency_score']):
            continue
        row['explanation'] = raw.get(explanation_column) or ''
        if 'score_source' in raw:
            row['score_source'] = raw['score_source']
        if raw.get('story_sha256'):
            row['story_sha256'] = raw['story_sha256']
        elif raw.get('response_text'):
            row['story_sha256'] = story_hash(raw['response_text'])
        result.append(validate_scores(row))
    return result


def check_score_stories(scores, stories):
    """Match IDs and, when available, the exact text actually annotated."""
    stories = {r['response_id']: r for r in stories}
    for row in scores:
        rid = row['response_id']
        if rid not in stories:
            raise ValueError('Score ID absent from datasets: ' + str(rid))
        if row.get('story_sha256') and row['story_sha256'] != story_hash(stories[rid]['response_text']):
            raise ValueError('Score story text hash mismatch: ' + rid)


def add_score_arguments(parser):
    parser.add_argument('--sheet')
    parser.add_argument('--id-column', default='response_id')
    parser.add_argument('--component-columns', nargs=3,
                        metavar=('INITIATIVE', 'RESOLUTION', 'RECOGNITION'))
    parser.add_argument('--total-column')
    parser.add_argument('--explanation-column')


def score_options(args):
    return {key: getattr(args, key) for key in
            ['sheet', 'id_column', 'component_columns', 'total_column', 'explanation_column']}
